import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_agile_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Agile Template"
    ws.views.sheetView[0].showGridLines = True
    
    # ---------------------------------------------------------
    # Styles matching the user's reference image
    # ---------------------------------------------------------
    # Header fill: Warm peach / salmon tone as seen in screenshot
    header_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="000000")
    
    data_font = Font(name="Calibri", size=10, color="000000")
    
    # Border
    thin_border_side = Side(border_style="thin", color="000000")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    # Exact 8 columns as requested
    headers = [
        "Planned Sprint",
        "Actual Sprint",
        "US ID",
        "User Story Description",
        "MOSCOW",
        "Dependency",
        "Assignee",
        "Status"
    ]
    
    ws.append(headers)
    ws.row_dimensions[1].height = 50
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
        
    # ---------------------------------------------------------
    # User Stories Data customized for Udhaya Kiran's Finora AI project
    # ---------------------------------------------------------
    stories = [
        [
            1,
            1,
            "US/101",
            "Collect real company financial documents (annual and quarterly reports, financial statements, SEC/regulatory filings, earnings reports, investor presentations, MD&A, notes to financial statements) for approx. 5-10 companies, with 2-3 years of annual reports each (suggested: Apple, Microsoft, Tesla, Amazon, Reliance Industries, Infosys, TCS; FY2023-2025)",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            1,
            1,
            "US/102",
            "Prepare the 10 assigned LLM topics (LLM basics, tokens, context window, Transformer, self-attention, prompting, temperature, hallucinations, next-token prediction, LLM limitations and mitigations such as RAG, citations, structured outputs, tool calling, validation and grounding) for the Q&A session beginning 10 Sep",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            2,
            2,
            "US/103",
            "Study financial document structures (annual reports, 10-K filings, earnings transcripts) and finalize the multi-agent architecture: orchestration flow, agent responsibilities and data models",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            2,
            2,
            "US/104",
            "Develop the research workspace and session management: create named research sessions, manage uploaded documents and organize research work",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            2,
            2,
            "US/105",
            "Build the Document Agent: document upload, PDF parsing, chunking, embedding generation, vector database indexing and 3-4 preloaded seed company documents",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            2,
            2,
            "US/106",
            "Build the Extraction Agent to automatically extract financial metrics, financial ratios, revenue trends and key performance figures",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            3,
            3,
            "US/107",
            "Build the Red Flag Agent to scan for anomalies, rising debt levels, falling margins, auditor qualifications and unusual financial patterns",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            3,
            3,
            "US/108",
            "Implement the multi-agent orchestration layer that coordinates the Document, Extraction and Red Flag agents sequentially on document upload",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            3,
            3,
            "US/109",
            "Test and validate agent outputs against seed financial documents for accuracy, source faithfulness and correctness",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            3,
            3,
            "US/110",
            "Build the Research Agent to handle multi-part queries, retrieve information step by step and give reasoned answers with strict source citations",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            3,
            3,
            "US/111",
            "Build the Comparison Agent to cross-reference extracted financial information and compare multiple company documents side by side",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            3,
            3,
            "US/112",
            "Develop a clean chat-style conversational research interface that returns cited, reasoned answers to financial questions",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            3,
            3,
            "US/113",
            "Begin Report Agent development: define the report structure and generate the Executive Summary and Key Financials sections",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            4,
            4,
            "US/114",
            "Complete the Report Agent to compile agent outputs into a structured PDF report (Executive Summary, Key Financials, Red Flags, Company Comparison, Outlook)",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            4,
            4,
            "US/115",
            "Conduct end-to-end testing: citation accuracy, agent collaboration, financial extraction correctness and report completeness",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ],
        [
            4,
            4,
            "US/116",
            "Prepare a project presentation and explain Finora AI (architecture, agents and demo) in the mock assessment on 15 Oct 2026",
            "MUST HAVE",
            "None",
            "Udhaya Kiran",
            "3- Completed"
        ]
    ]
    
    col_widths = {
        "A": 14,   # Planned Sprint
        "B": 14,   # Actual Sprint
        "C": 12,   # US ID
        "D": 72,   # User Story Description
        "E": 16,   # MOSCOW
        "F": 16,   # Dependency
        "G": 18,   # Assignee
        "H": 16    # Status
    }
    
    alignments = {
        "A": Alignment(horizontal="center", vertical="center"),
        "B": Alignment(horizontal="center", vertical="center"),
        "C": Alignment(horizontal="center", vertical="center"),
        "D": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "E": Alignment(horizontal="center", vertical="center"),
        "F": Alignment(horizontal="center", vertical="center"),
        "G": Alignment(horizontal="center", vertical="center"),
        "H": Alignment(horizontal="center", vertical="center")
    }
    
    row_heights = {
        2: 55,   # US/101
        3: 55,   # US/102
        4: 45,   # US/103
        5: 45,   # US/104
        6: 45,   # US/105
        7: 45,   # US/106
        8: 45,   # US/107
        9: 45,   # US/108
        10: 45,  # US/109
        11: 45,  # US/110
        12: 45,  # US/111
        13: 45,  # US/112
        14: 45,  # US/113
        15: 45,  # US/114
        16: 45,  # US/115
        17: 45   # US/116
    }
    
    for row in stories:
        ws.append(row)
        
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = row_heights.get(row_idx, 45)
        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = cell_border
            cell.alignment = alignments.get(col_letter, Alignment(horizontal="center", vertical="center"))
            
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
        
    output_path = os.path.join(os.path.dirname(__file__), "Agile_Template.xlsx")
    wb.save(output_path)
    print(f"Agile template generated successfully at: {output_path}")

if __name__ == "__main__":
    build_agile_template()
