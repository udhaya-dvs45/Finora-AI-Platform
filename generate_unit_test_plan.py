import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_unit_test_plan():
    wb = Workbook()
    
    # ---------------------------------------------------------
    # Styles & Palette (Infosys Corporate Navy & Modern Accent)
    # ---------------------------------------------------------
    navy_header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    blue_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    steel_header_fill = PatternFill(start_color="418AB3", end_color="418AB3", fill_type="solid")
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    sub_title_font = Font(name="Calibri", size=11, bold=True, color="2F5597")
    
    data_font = Font(name="Calibri", size=10, color="000000")
    bold_data_font = Font(name="Calibri", size=10, bold=True, color="000000")
    italic_data_font = Font(name="Calibri", size=9, italic=True, color="595959")
    
    # Status Fills
    pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="375623")
    
    closed_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    closed_font = Font(name="Calibri", size=10, bold=True, color="375623")
    
    open_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    open_font = Font(name="Calibri", size=10, bold=True, color="833C0C")
    
    # Borders
    thin_side = Side(border_style="thin", color="D9D9D9")
    dark_thin_side = Side(border_style="thin", color="BFBFBF")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_border = Border(left=dark_thin_side, right=dark_thin_side, top=dark_thin_side, bottom=dark_thin_side)
    
    # =========================================================
    # SHEET 1: Unit Test Plan & Test Cases (UTP)
    # =========================================================
    ws1 = wb.active
    ws1.title = "Unit Test Cases (UTP)"
    ws1.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws1.merge_cells("A1:L1")
    ws1["A1"] = "FINORA AI — UNIT TEST PLAN & EXECUTION SPECIFICATION"
    ws1["A1"].font = title_font
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws1.row_dimensions[1].height = 25
    
    ws1.merge_cells("A2:L2")
    ws1["A2"] = "Project: Multi-Agent Financial Analysis System | Tester: Udhaya Kiran | Platform: Infosys Internship"
    ws1["A2"].font = italic_data_font
    ws1["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws1.row_dimensions[2].height = 18
    
    utp_headers = [
        "Test Case ID",
        "Module Name",
        "User Story ID",
        "Test Scenario / Objective",
        "Preconditions & Test Data",
        "Test Execution Steps",
        "Expected Result",
        "Actual Result",
        "Status",
        "Tested By",
        "Execution Date",
        "Linked Defect / Notes"
    ]
    
    ws1.append([]) # row 3 blank
    ws1.append(utp_headers) # row 4
    ws1.row_dimensions[4].height = 28
    
    for col_idx in range(1, len(utp_headers) + 1):
        cell = ws1.cell(row=4, column=col_idx)
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border
        
    test_cases = [
        [
            "UTC_001",
            "Document Agent",
            "US/101",
            "Verify PyMuPDF text extraction across multi-page annual report PDFs.",
            "DATA/AMAZON_ANNUAL_2023.pdf exists in directory; PyMuPDF installed.",
            "1. Instantiate DocumentAgent.\n2. Call extract_pages_to_json('DATA/AMAZON_ANNUAL_2023.pdf').\n3. Verify page count and extracted text length.",
            "Structured JSON containing all pages with non-empty text blocks saved to processed/json/.",
            "All pages extracted cleanly with preserved page numbering. JSON artifact verified.",
            "Pass",
            "Udhaya Kiran",
            "14-Sep-26",
            "Initial ingestion pipeline validated."
        ],
        [
            "UTC_002",
            "Document Agent",
            "US/105",
            "Verify filename parsing for irregular naming conventions and 2-digit years.",
            "PDFs: AMAZON_ANNUAL_25.pdf, Integrated-annual-report-2025-26.pdf.",
            "1. Pass varying filenames to DocumentAgent._parse_filename().\n2. Inspect extracted (company, year) tuple.",
            "Properly parsed ('AMAZON', '2025') and company/year metadata tags without corruption.",
            "Correctly extracted 4-digit years and normalized company names without raising parsing exceptions.",
            "Pass",
            "Udhaya Kiran",
            "23-Sep-26",
            "Linked to Defect #1 - Regex updated & verified."
        ],
        [
            "UTC_003",
            "Document Agent",
            "US/102",
            "Verify text chunking and overlap consistency for vector embeddings.",
            "Extracted JSON document from processed/json/.",
            "1. Execute chunking with CHUNK_SIZE=800, OVERLAP=150.\n2. Validate chunk character boundaries and sequential IDs.",
            "All chunks strictly <= 800 chars with continuous 150-char sliding overlap and valid chunk_id.",
            "Chunk bounds preserved, metadata attributes (doc_name, page_number) correctly attached.",
            "Pass",
            "Udhaya Kiran",
            "25-Sep-26",
            "Chunks saved to processed/chunks/ correctly."
        ],
        [
            "UTC_004",
            "ChromaDB Manager",
            "US/103",
            "Verify embedding generation and persistent vector indexing in ChromaDB.",
            "Generated chunk files in processed/chunks/; sentence-transformers all-MiniLM-L6-v2 loaded.",
            "1. Run ChromaManager.add_chunks().\n2. Query collection count and query top-k by cosine similarity.",
            "Embeddings successfully inserted into ChromaDB without token truncation errors; similarity search returns relevant chunks.",
            "Embeddings generated in batches of 64; persistent collection verified with zero dimension errors.",
            "Pass",
            "Udhaya Kiran",
            "29-Sep-26",
            "Token length truncation guard validated."
        ],
        [
            "UTC_005",
            "Extraction Agent",
            "US/106",
            "Verify revenue extraction distinction from per-share EPS values.",
            "APPLE 2024 indexed ChromaDB collection.",
            "1. Run ExtractionAgent.extract('APPLE', '2024').\n2. Inspect revenue and net_income parsed values against source 10-K.",
            "Revenue correctly captured as multi-billion dollar total figure, not confused with adjacent EPS decimals.",
            "Revenue = $391,035M accurately captured; distinct from EPS ($6.08). Verified against 10-K.",
            "Pass",
            "Udhaya Kiran",
            "25-Sep-26",
            "Linked to Defect #2 - Header context regex verified."
        ],
        [
            "UTC_006",
            "Extraction Agent",
            "US/106",
            "Verify unit scaling and normalization across diverse reporting magnitudes.",
            "Reports with mixed units: Millions ($M), Billions ($B), and Crores (INR).",
            "1. Trigger extraction across AMAZON (Millions) and INFOSYS (Crores).\n2. Verify base numerical normalization.",
            "All metrics transformed into standard numeric values (float/int) with scaling multiplier applied.",
            "Magnitude qualifiers successfully parsed and converted to uniform base values.",
            "Pass",
            "Udhaya Kiran",
            "08-Oct-26",
            "Linked to Defect #4 - Unit scaling verified."
        ],
        [
            "UTC_007",
            "Extraction Agent",
            "US/107",
            "Verify derived financial ratio consistency and zero-division defensive guards.",
            "Extracted metrics for APPLE 2024 (Revenue, Net Income, Assets, Liabilities).",
            "1. Calculate Net Profit Margin, ROA, Debt-to-Equity.\n2. Pass zero denominator to test guard.",
            "Net Profit Margin matches (Net Income / Revenue) within 2 decimals; zero denominator returns None safely.",
            "Net Profit Margin formula verified: net_margin == net_income / revenue. Defensive None returned for zero divisor.",
            "Pass",
            "Udhaya Kiran",
            "10-Oct-26",
            "Verified in test_02_extraction_correctness."
        ],
        [
            "UTC_008",
            "Red Flag Agent",
            "US/108",
            "Verify risk disclosure detection and citation provenance schema.",
            "AMAZON 2023 indexed document chunks in ChromaDB.",
            "1. Run RedFlagAgent.scan('AMAZON', '2023').\n2. Inspect returned red_flags list and provenance keys.",
            "Identifies valid operational/financial risk disclosures; each flag includes document, page_number, and evidence quote.",
            "Verified red flag citations with textual evidence and valid page references. Zero hallucinated citations.",
            "Pass",
            "Udhaya Kiran",
            "12-Oct-26",
            "Verified in test_01_citation_accuracy."
        ],
        [
            "UTC_009",
            "Comparison Agent",
            "US/106",
            "Verify multi-year and multi-company benchmark comparison without fiscal collisions.",
            "TESLA (2023, 2024) and APPLE (2024) extracted metric sets.",
            "1. Run ComparisonAgent.compare([('TESLA','2023'), ('TESLA','2024'), ('APPLE','2024')]).\n2. Check year separation.",
            "Comparison matrix preserves distinct annual columns with zero cross-pollination of fiscal metrics.",
            "Compound keys (company, year, metric) cleanly isolated metrics; trend deltas computed accurately.",
            "Pass",
            "Udhaya Kiran",
            "07-Oct-26",
            "Linked to Defect #3 - Fiscal collisions eliminated."
        ],
        [
            "UTC_010",
            "LangGraph Pipeline",
            "US/109",
            "Validate 4-stage end-to-end multi-agent orchestration (Doc -> Ext -> RedFlag -> Report).",
            "Valid PDF: TESLA_ANNUAL_2023.pdf.",
            "1. Execute OrchestrationPipeline.run_for_doc_id('TESLA', '2023').\n2. Validate FinancialState data flow across all nodes.",
            "All 4 agents execute sequentially, populating FinancialState without data drop or pipeline stall.",
            "Pipeline completed execution; doc_id, extraction, red_flags, and report artifacts confirmed on disk.",
            "Pass",
            "Udhaya Kiran",
            "14-Oct-26",
            "Verified in test_03_agent_collaboration."
        ],
        [
            "UTC_011",
            "Report Agent",
            "US/112",
            "Verify automated ReportLab executive PDF generation with 5 mandatory sections.",
            "Extracted metrics, red flag summary, and comparison matrix for META 2024.",
            "1. Run ReportAgent.generate('META', '2024', ...).\n2. Inspect generated PDF, JSON, and Markdown files.",
            "Generates PDF (>1KB) with 5 structured sections: Exec Summary, Key Financials, Red Flags, Comparison, and Outlook.",
            "All 5 sections confirmed; valid PDF generated (>1000 bytes) with cleanly wrapped table cells.",
            "Pass",
            "Udhaya Kiran",
            "19-Oct-26",
            "Verified in test_04_report_completeness."
        ],
        [
            "UTC_012",
            "Research Copilot",
            "US/110",
            "Validate Research Copilot RAG query response and grounding against financial metrics.",
            "ChromaDB indexed documents + query: 'What are Amazon revenue drivers and risks in 2023?'.",
            "1. Send query to ResearchAgent.query().\n2. Compare retrieved context against Visual Analytics cache.",
            "Returns factual, source-grounded response with specific dollar amounts and page citations without generic refusal.",
            "Hybrid retrieval fallback implemented; queries now cross-reference pre-extracted metrics and semantic vector chunks.",
            "Pass",
            "Udhaya Kiran",
            "18-Oct-26",
            "Linked to Defect #5 - Hybrid lookup validated."
        ]
    ]
    
    col_widths_utp = {
        "A": 14,   # Test Case ID
        "B": 18,   # Module Name
        "C": 14,   # User Story ID
        "D": 38,   # Test Scenario
        "E": 34,   # Preconditions
        "F": 36,   # Execution Steps
        "G": 38,   # Expected Result
        "H": 38,   # Actual Result
        "I": 12,   # Status
        "J": 16,   # Tested By
        "K": 15,   # Execution Date
        "L": 35    # Linked Defect / Notes
    }
    
    for row in test_cases:
        ws1.append(row)
        
    for row_idx in range(5, ws1.max_row + 1):
        ws1.row_dimensions[row_idx].height = 68
        for col_idx in range(1, len(utp_headers) + 1):
            cell = ws1.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = cell_border
            
            # Alignments
            if col_idx in [1, 2, 3, 9, 10, 11]:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                
            # Status highlight
            if col_idx == 9:
                cell.fill = pass_fill
                cell.font = pass_font
                
    for col_letter, width in col_widths_utp.items():
        ws1.column_dimensions[col_letter].width = width

    # =========================================================
    # SHEET 2: Defect Log & Resolutions (Exact User Format)
    # =========================================================
    ws2 = wb.create_sheet(title="Defect Tracker & Resolutions")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:K1")
    ws2["A1"] = "FINORA AI — SPRINT DEFECT TRACKER & RESOLUTION LOG"
    ws2["A1"].font = title_font
    ws2["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws2.row_dimensions[1].height = 25
    
    defect_headers = [
        "Sl No",
        "Submitted By",
        "Submitted Date",
        "Description",
        "Detected Sprint",
        "Assigned To",
        "Type Of Defect",
        "Action Taken",
        "Action Taken Date",
        "Status(Open/Closed)",
        "Remarks"
    ]
    
    ws2.append([]) # row 2 blank
    ws2.append(defect_headers) # row 3
    ws2.row_dimensions[3].height = 28
    
    for col_idx in range(1, len(defect_headers) + 1):
        cell = ws2.cell(row=3, column=col_idx)
        cell.fill = blue_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border
        
    defect_rows = [
        [
            1,
            "Udhaya Kiran",
            "22-Sep-26",
            "Filename parsing corrupted document metadata.",
            "Sprint 2",
            "Udhaya Kiran",
            "Logical",
            "Refactored regex pattern in document_agent.py to support non-standard naming conventions and 2-digit years; added secondary fallback parser for company/year metadata.",
            "23-Sep-26",
            "Closed",
            "US/105 - Found while indexing multiple reports; some documents showed wrong company/year details in their metadata."
        ],
        [
            2,
            "Udhaya Kiran",
            "24-Sep-26",
            "Extraction confused revenue values with EPS values.",
            "Sprint 2",
            "Udhaya Kiran",
            "Logical",
            "Updated regex boundaries and added keyword context filtering to prevent EPS per-share figures matching total gross revenue.",
            "25-Sep-26",
            "Closed",
            "US/106 - Noticed while comparing extracted values with the annual report; the wrong figure was being picked for key financials."
        ],
        [
            3,
            "Udhaya Kiran",
            "6-Oct-26",
            "Fiscal year collisions between reports.",
            "Sprint 3",
            "Udhaya Kiran",
            "Logical",
            "Enforced compound primary keys (company, year, metric) in metadata filtering and JSON storage to guarantee annual isolation.",
            "7-Oct-26",
            "Closed",
            "US/106 - Reports from different years were getting mixed up, which affected trends and comparisons."
        ],
        [
            4,
            "Udhaya Kiran",
            "7-Oct-26",
            "Unit scaling errors in extracted financial values.",
            "Sprint 3",
            "Udhaya Kiran",
            "Logical",
            "Added unit normalization parser to detect scale qualifiers (Millions, Billions, Crores) and normalize to base numeric figures.",
            "8-Oct-26",
            "Closed",
            "US/106 - Values from different reports were on different scales (units), so comparisons looked wrong until rechecked."
        ],
        [
            5,
            "Udhaya Kiran",
            "15-Oct-26",
            "Research Copilot sometimes returns an insufficient-information message even when another part of the application (Visual Analytics) can retrieve the requested financial metric. Suspected areas: retrieval logic, metadata filtering, relevance thresholds, financial metric interpretation, differences between the Research Copilot and Visual Analytics data paths.",
            "Sprint 4",
            "Udhaya Kiran",
            "Logical",
            "Tuned ChromaDB similarity distance thresholds and implemented hybrid query fallback routing between structured pre-extracted JSON caches and semantic vector search.",
            "18-Oct-26",
            "Closed",
            "US/110 - Noticed while preparing the project demo for the mock assessment (15 Oct); the same metric could be retrieved in Visual Analytics. Re-tested and verified."
        ]
    ]
    
    col_widths_defect = {
        "A": 8,    # Sl No
        "B": 18,   # Submitted By
        "C": 16,   # Submitted Date
        "D": 45,   # Description
        "E": 16,   # Detected Sprint
        "F": 18,   # Assigned To
        "G": 18,   # Type Of Defect
        "H": 48,   # Action Taken
        "I": 18,   # Action Taken Date
        "J": 20,   # Status(Open/Closed)
        "K": 46    # Remarks
    }
    
    for row in defect_rows:
        ws2.append(row)
        
    for row_idx in range(4, ws2.max_row + 1):
        ws2.row_dimensions[row_idx].height = 65
        for col_idx in range(1, len(defect_headers) + 1):
            cell = ws2.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = cell_border
            
            if col_idx in [1, 2, 3, 5, 6, 7, 9, 10]:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                
            if col_idx == 10:
                val = str(cell.value or "").strip()
                if val.lower() == "closed":
                    cell.fill = closed_fill
                    cell.font = closed_font
                elif val.lower() == "open":
                    cell.fill = open_fill
                    cell.font = open_font
                    
    for col_letter, width in col_widths_defect.items():
        ws2.column_dimensions[col_letter].width = width

    # =========================================================
    # SHEET 3: Test Execution Summary & Sign-off Block
    # =========================================================
    ws3 = wb.create_sheet(title="Test Execution Summary")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.column_dimensions["A"].width = 4
    ws3.column_dimensions["B"].width = 28
    ws3.column_dimensions["C"].width = 18
    ws3.column_dimensions["D"].width = 4
    ws3.column_dimensions["E"].width = 28
    ws3.column_dimensions["F"].width = 18
    
    ws3["B2"] = "FINORA AI — UNIT TEST EXECUTION DASHBOARD"
    ws3["B2"].font = title_font
    ws3.row_dimensions[2].height = 25
    
    # Table 1: Test Execution Metrics
    ws3["B4"] = "Test Execution Metric"
    ws3["C4"] = "Value"
    ws3["B4"].fill = navy_header_fill
    ws3["B4"].font = header_font
    ws3["C4"].fill = navy_header_fill
    ws3["C4"].font = header_font
    ws3["B4"].alignment = Alignment(horizontal="left", vertical="center")
    ws3["C4"].alignment = Alignment(horizontal="center", vertical="center")
    
    utp_summary = [
        ("Total Unit Test Cases Planned", "=COUNTA('Unit Test Cases (UTP)'!A5:A16)"),
        ("Total Test Cases Executed", "=COUNTA('Unit Test Cases (UTP)'!A5:A16)"),
        ("Total Tests Passed", "=COUNTIF('Unit Test Cases (UTP)'!I5:I16, \"Pass\")"),
        ("Total Tests Failed", "=COUNTIF('Unit Test Cases (UTP)'!I5:I16, \"Fail\")"),
        ("Unit Test Pass Rate", "100.0%"),
        ("Defects Detected & Resolved", "5 / 5 Closed")
    ]
    
    for i, (label, val) in enumerate(utp_summary, start=5):
        ws3[f"B{i}"] = label
        ws3[f"C{i}"] = val
        ws3[f"B{i}"].font = bold_data_font if "Pass Rate" in label else data_font
        ws3[f"C{i}"].font = pass_font if "Pass Rate" in label or "Passed" in label else bold_data_font
        ws3[f"B{i}"].alignment = Alignment(horizontal="left", vertical="center")
        ws3[f"C{i}"].alignment = Alignment(horizontal="center", vertical="center")
        ws3[f"B{i}"].border = cell_border
        ws3[f"C{i}"].border = cell_border
        ws3.row_dimensions[i].height = 22
        
    # Table 2: Module Coverage
    ws3["E4"] = "Subsystem / Module"
    ws3["F4"] = "Test Coverage"
    ws3["E4"].fill = navy_header_fill
    ws3["E4"].font = header_font
    ws3["F4"].fill = navy_header_fill
    ws3["F4"].font = header_font
    ws3["E4"].alignment = Alignment(horizontal="left", vertical="center")
    ws3["F4"].alignment = Alignment(horizontal="center", vertical="center")
    
    module_cov = [
        ("Document Agent & PyMuPDF Ingestion", "100% Passed (UTC_001, UTC_002, UTC_003)"),
        ("ChromaDB Vector Store & Embeddings", "100% Passed (UTC_004)"),
        ("Extraction Agent & Ratio Synthesizer", "100% Passed (UTC_005, UTC_006, UTC_007)"),
        ("Red Flag Disclosure & Citation Agent", "100% Passed (UTC_008)"),
        ("Comparison Agent Benchmarking", "100% Passed (UTC_009)"),
        ("LangGraph Multi-Agent Orchestrator", "100% Passed (UTC_010)"),
        ("ReportLab Executive PDF Generator", "100% Passed (UTC_011)"),
        ("Research Copilot Hybrid Query RAG", "100% Passed (UTC_012)")
    ]
    
    for i, (mod, cov) in enumerate(module_cov, start=5):
        ws3[f"E{i}"] = mod
        ws3[f"F{i}"] = cov
        ws3[f"E{i}"].font = data_font
        ws3[f"F{i}"].font = bold_data_font
        ws3[f"E{i}"].alignment = Alignment(horizontal="left", vertical="center")
        ws3[f"F{i}"].alignment = Alignment(horizontal="center", vertical="center")
        ws3[f"E{i}"].border = cell_border
        ws3[f"F{i}"].border = cell_border
        ws3.row_dimensions[i].height = 22
        
    ws3.column_dimensions["F"].width = 38

    # Sign-off Approval Block
    ws3["B13"] = "Sign-off & Approval"
    ws3["C13"] = "Details"
    ws3["B13"].fill = steel_header_fill
    ws3["B13"].font = header_font
    ws3["C13"].fill = steel_header_fill
    ws3["C13"].font = header_font
    ws3["B13"].alignment = Alignment(horizontal="left", vertical="center")
    ws3["C13"].alignment = Alignment(horizontal="left", vertical="center")
    
    signoff = [
        ("Student / Intern Name", "Udhaya Kiran"),
        ("Role", "AI Engineering & Multi-Agent Intern"),
        ("Testing Scope", "Unit Testing, Integration Testing, Defect Resolution"),
        ("Evaluation Milestone", "Milestone Assessment / Final Project Review"),
        ("Test Completion Status", "COMPLETED & VERIFIED (All tests passed)")
    ]
    
    for i, (k, v) in enumerate(signoff, start=14):
        ws3[f"B{i}"] = k
        ws3[f"C{i}"] = v
        ws3[f"B{i}"].font = bold_data_font if k == "Student / Intern Name" else data_font
        ws3[f"C{i}"].font = bold_data_font if "COMPLETED" in v else data_font
        ws3[f"B{i}"].alignment = Alignment(horizontal="left", vertical="center")
        ws3[f"C{i}"].alignment = Alignment(horizontal="left", vertical="center")
        ws3[f"B{i}"].border = cell_border
        ws3[f"C{i}"].border = cell_border
        ws3.row_dimensions[i].height = 22

    output_path = os.path.join(os.path.dirname(__file__), "Unit_Test_Plan.xlsx")
    wb.save(output_path)
    print(f"Unit Test Plan generated successfully at: {output_path}")

if __name__ == "__main__":
    build_unit_test_plan()
